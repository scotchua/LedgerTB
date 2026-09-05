"""The file-export seam to workpaper tools (LedgerPDF pairing).

Contract: exports only land inside user-approved roots, work at read level
(with the action audit-logged), and the Excel carries computed totals — not
uncalculated =SUM() formulas that render as literal text in any tool without
a calc engine.
"""
import os
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
import pypdfium2 as pdfium

from database import connection as dbconn
from database.connection import get_connection
from models.audit_log import AuditLog
from models.client import Client
from services import mcp_tools
from services.close_package import (
    SNAPSHOT_CANONICALIZATION_VERSION,
    close_package_snapshot_hash,
    close_package_snapshot_payload,
    load_close_package_snapshot,
)
from tests.conftest import post_entry


def _seed(client_id, accounts):
    post_entry(client_id, date(2026, 1, 15),
               [(accounts["cash"], 500, 0), (accounts["revenue"], 0, 500)])
    post_entry(client_id, date(2026, 2, 3),
               [(accounts["expense"], 120, 0), (accounts["cash"], 0, 120)])


def _pdf_text(payload):
    document = pdfium.PdfDocument(payload)
    try:
        return "\n".join(
            document[index].get_textpage().get_text_range()
            for index in range(len(document))
        )
    finally:
        document.close()


def _export_state(client_id, directory):
    conn = get_connection()
    try:
        audit = conn.execute(
            "SELECT id FROM document_audits WHERE client_id = ?",
            (client_id,),
        ).fetchone()
        completion = conn.execute(
            "SELECT id FROM audit_log WHERE client_id = ? "
            "AND table_name = 'close_package_issued'",
            (client_id,),
        ).fetchone()
    finally:
        conn.close()
    final_files = list(directory.glob("*.pdf")) + list(directory.glob("*.xlsx"))
    return audit, completion, final_files


def test_export_refused_without_roots(client_id, accounts, tmp_path, monkeypatch):
    _seed(client_id, accounts)
    monkeypatch.delenv("LEDGERTB_MCP_EXPORT_ROOTS", raising=False)
    monkeypatch.delenv("PROBOOKS_MCP_EXPORT_ROOTS", raising=False)
    with pytest.raises(ValueError, match="export is off"):
        mcp_tools.export_close_package(client_id, "2026-01-01", "2026-03-31",
                                       str(tmp_path))


def test_export_refused_outside_roots(client_id, accounts, tmp_path, monkeypatch):
    _seed(client_id, accounts)
    approved = tmp_path / "approved"
    approved.mkdir()
    elsewhere = tmp_path / "elsewhere"
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(approved))
    with pytest.raises(ValueError, match="outside") as exc_info:
        mcp_tools.export_close_package(client_id, "2026-01-01", "2026-03-31",
                                       str(elsewhere))
    assert str(approved) in str(exc_info.value)
    with pytest.raises(ValueError, match="outside"):
        mcp_tools.export_close_package(client_id, "2026-01-01", "2026-03-31",
                                       str(approved / ".." / "elsewhere"))


def test_export_writes_both_files_at_read_level(client_id, accounts, tmp_path,
                                                monkeypatch):
    _seed(client_id, accounts)
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(tmp_path))
    monkeypatch.setattr(dbconn, "ASSISTANT_ACCESS_LEVEL", "read")

    result = mcp_tools.export_close_package(
        client_id, "2026-01-01", "2026-03-31", str(tmp_path / "binder-src"))

    pdf = open(result["pdf"], "rb").read()
    assert pdf.startswith(b"%PDF")

    wb = openpyxl.load_workbook(BytesIO(open(result["xlsx"], "rb").read()))
    assert set(wb.sheetnames) >= {
        "Summary", "Income Statement", "Balance Sheet", "Trial Balance",
        "Transactions",
    }

    # Gap-2 regression: the TOTALS row must be numbers, never formula text —
    # openpyxl caches no results, so "=SUM(...)" reads as a literal string in
    # any tool without a calc engine.
    tb = wb["Trial Balance"]
    totals_row = tb.max_row
    assert tb.cell(row=totals_row, column=1).value == "TOTALS"
    for col in range(4, 12):
        value = tb.cell(row=totals_row, column=col).value
        assert not (isinstance(value, str) and value.startswith("=")), \
            f"col {col} is an uncalculated formula"
        assert isinstance(value, (int, float))
    assert tb.cell(row=totals_row, column=10).value == pytest.approx(500.0)

    start = date(2026, 1, 1)
    end = date(2026, 3, 31)
    client = Client.get_by_id(client_id)
    tb_rows, _ = mcp_tools.ReportGenerator.trial_balance_worksheet(
        client_id, start, end
    )
    snapshot = load_close_package_snapshot(client_id, start, end)
    expected_hash = close_package_snapshot_hash(close_package_snapshot_payload(
        client_id, client.name, start, end, tb_rows, snapshot
    ))
    conn = get_connection()
    try:
        audit = conn.execute(
            "SELECT id, doc_key, content_hash, canonicalization_version "
            "FROM document_audits WHERE client_id = ?",
            (client_id,),
        ).fetchone()
        audit_log = conn.execute(
            "SELECT record_id FROM audit_log "
            "WHERE client_id = ? AND table_name = 'document_audits' "
            "AND action = 'EXPORT'",
            (client_id,),
        ).fetchone()
    finally:
        conn.close()
    assert audit is not None
    assert audit["doc_key"] == f"{client_id}:2026-01-01:2026-03-31"
    assert audit["content_hash"] == expected_hash
    assert audit["canonicalization_version"] == SNAPSHOT_CANONICALIZATION_VERSION
    assert audit_log is not None
    assert audit_log["record_id"] == audit["id"]
    conn = get_connection()
    try:
        completion = conn.execute(
            "SELECT record_id FROM audit_log "
            "WHERE client_id = ? AND table_name = 'close_package_issued'",
            (client_id,),
        ).fetchone()
    finally:
        conn.close()
    assert completion is not None
    assert completion["record_id"] == audit["id"]

    text = _pdf_text(pdf)
    assert f"Document Audits, ID {audit['id']}" in text
    assert "Document Audits, ID None" not in text

    # The export is audit-logged even at read level.
    monkeypatch.setattr(dbconn, "ASSISTANT_ACCESS_LEVEL", None)
    counts = AuditLog.get_filtered_counts(client_id)
    assert counts["total"] > 0


def test_failure_after_audit_reservation_leaves_no_issued_files(
    client_id, accounts, tmp_path, monkeypatch
):
    _seed(client_id, accounts)
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(tmp_path))
    target = tmp_path / "binder-src"

    def fail_before_temp_write(*args, **kwargs):
        raise RuntimeError("before file rename")

    monkeypatch.setattr(mcp_tools, "_write_private_temp", fail_before_temp_write)
    with pytest.raises(RuntimeError, match="before file rename"):
        mcp_tools.export_close_package(
            client_id, "2026-01-01", "2026-03-31", str(target)
        )

    audit, completion, final_files = _export_state(client_id, target)
    assert audit is not None
    assert completion is None
    assert final_files == []


def test_render_failure_leaves_no_final_or_temporary_files(
    client_id, accounts, tmp_path, monkeypatch
):
    _seed(client_id, accounts)
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(tmp_path))
    target = tmp_path / "binder-src"

    def failed_render(*args, **kwargs):
        raise RuntimeError("render failed")

    monkeypatch.setattr(
        "services.close_package.SimpleDocTemplate.build", failed_render
    )
    with pytest.raises(RuntimeError, match="render failed"):
        mcp_tools.export_close_package(
            client_id, "2026-01-01", "2026-03-31", str(target)
        )

    audit, completion, final_files = _export_state(client_id, target)
    assert audit is not None
    assert completion is None
    assert final_files == []
    assert list(target.glob(".*.tmp")) == []


def test_second_rename_failure_removes_first_final_file(
    client_id, accounts, tmp_path, monkeypatch
):
    _seed(client_id, accounts)
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(tmp_path))
    target = tmp_path / "binder-src"
    original_replace = Path.replace

    def fail_xlsx_rename(path, destination):
        if Path(destination).suffix == ".xlsx":
            raise RuntimeError("xlsx rename failed")
        return original_replace(path, destination)

    monkeypatch.setattr(Path, "replace", fail_xlsx_rename)
    with pytest.raises(RuntimeError, match="xlsx rename failed"):
        mcp_tools.export_close_package(
            client_id, "2026-01-01", "2026-03-31", str(target)
        )

    audit, completion, final_files = _export_state(client_id, target)
    assert audit is not None
    assert completion is None
    assert final_files == []


def test_export_defaults_to_the_approved_folder(client_id, accounts, tmp_path,
                                                monkeypatch):
    _seed(client_id, accounts)
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(tmp_path))

    result = mcp_tools.export_close_package(
        client_id, "2026-01-01", "2026-03-31", ""
    )

    assert Path(result["pdf"]).parent == tmp_path
    assert Path(result["xlsx"]).parent == tmp_path


def test_relative_export_directory_is_resolved_inside_approved_root(
    client_id, accounts, tmp_path, monkeypatch
):
    _seed(client_id, accounts)
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(tmp_path))

    result = mcp_tools.export_close_package(
        client_id, "2026-01-01", "2026-03-31", "  Q1 close  "
    )

    assert Path(result["pdf"]).parent == tmp_path / "Q1 close"


def test_the_folder_chosen_in_the_app_beats_the_environment(client_id, accounts,
                                                            tmp_path,
                                                            monkeypatch):
    """The user's in-app consent must win. An MCP server's environment comes
    from the client's config file, which a shell-capable assistant can edit —
    if the env var overrode the vault, the assistant could rewrite the very
    boundary the user set for it."""
    from utils.secure_store import set_secret
    from utils.assistant_access import credential_names

    _seed(client_id, accounts)
    monkeypatch.delenv("LEDGERTB_MCP_EXPORT_ROOTS", raising=False)
    monkeypatch.delenv("PROBOOKS_MCP_EXPORT_ROOTS", raising=False)

    vault_root = tmp_path / "vault-root"
    vault_root.mkdir()
    names = credential_names(dbconn.DATABASE_PATH)
    set_secret(names.export_roots, str(vault_root))
    result = mcp_tools.export_close_package(
        client_id, "2026-01-01", "2026-03-31", str(vault_root / "julyco"))
    assert result["pdf"].startswith(str(vault_root))

    # An attacker-set environment root does not widen the boundary...
    env_root = tmp_path / "env-root"
    env_root.mkdir()
    monkeypatch.setenv("LEDGERTB_MCP_EXPORT_ROOTS", str(env_root))
    with pytest.raises(ValueError, match="outside"):
        mcp_tools.export_close_package(
            client_id, "2026-01-01", "2026-03-31", str(env_root / "escaped"))
    # ...and the chosen folder keeps working while it is set.
    again = mcp_tools.export_close_package(
        client_id, "2026-01-01", "2026-03-31", str(vault_root / "again"))
    assert again["pdf"].startswith(str(vault_root))


def test_export_refuses_to_write_through_a_planted_symlink(client_id, accounts,
                                                           tmp_path, monkeypatch):
    """The filename is predictable from client and period, so on a shared
    engagement folder a colleague could pre-place a symlink to catch the
    close package. Containment checks the directory; this checks the file."""
    from utils.secure_store import set_secret
    from utils.assistant_access import credential_names

    _seed(client_id, accounts)
    monkeypatch.delenv("LEDGERTB_MCP_EXPORT_ROOTS", raising=False)
    monkeypatch.delenv("PROBOOKS_MCP_EXPORT_ROOTS", raising=False)

    root = tmp_path / "engagement"
    root.mkdir()
    names = credential_names(dbconn.DATABASE_PATH)
    set_secret(names.export_roots, str(root))

    client = Client.get_by_id(client_id)
    stem = f"{client.name} close package 2026-01-01 to 2026-03-31"
    elsewhere = tmp_path / "colleague-copy.pdf"
    try:
        (root / f"{stem}.pdf").symlink_to(elsewhere)
    except (OSError, NotImplementedError):
        pytest.skip(
            "symlink creation not permitted on this platform without elevation"
        )

    with pytest.raises(ValueError, match="symbolic link"):
        mcp_tools.export_close_package(
            client_id, "2026-01-01", "2026-03-31", str(root))
    assert not elsewhere.exists(), "export was written through the symlink"


def test_symlink_export_check_skips_when_creation_is_not_permitted(
    client_id, accounts, tmp_path, monkeypatch
):
    def refuse_symlink(*args, **kwargs):
        raise OSError("symlink privilege not held")

    monkeypatch.setattr(os, "symlink", refuse_symlink)

    with pytest.raises(pytest.skip.Exception, match="without elevation"):
        test_export_refuses_to_write_through_a_planted_symlink(
            client_id, accounts, tmp_path, monkeypatch
        )


@pytest.mark.skipif(os.name == "nt",
                    reason="Windows has no POSIX file mode; NTFS inherits the "
                           "user-profile ACL instead")
def test_exports_are_not_world_readable(client_id, accounts, tmp_path,
                                        monkeypatch):
    """The book is 0600; the unencrypted close package of that same book
    must not be looser."""
    import stat

    from utils.secure_store import set_secret
    from utils.assistant_access import credential_names

    _seed(client_id, accounts)
    monkeypatch.delenv("LEDGERTB_MCP_EXPORT_ROOTS", raising=False)
    monkeypatch.delenv("PROBOOKS_MCP_EXPORT_ROOTS", raising=False)

    root = tmp_path / "exports"
    root.mkdir()
    names = credential_names(dbconn.DATABASE_PATH)
    set_secret(names.export_roots, str(root))

    result = mcp_tools.export_close_package(
        client_id, "2026-01-01", "2026-03-31", str(root / "q1"))
    for path in (result["pdf"], result["xlsx"]):
        mode = stat.S_IMODE(Path(path).stat().st_mode)
        assert mode == 0o600, f"{path} is {oct(mode)}"


def test_probooks_export_root_env_alias_remains_supported(client_id, accounts,
                                                          tmp_path,
                                                          monkeypatch):
    _seed(client_id, accounts)
    monkeypatch.delenv("LEDGERTB_MCP_EXPORT_ROOTS", raising=False)
    monkeypatch.setenv("PROBOOKS_MCP_EXPORT_ROOTS", str(tmp_path))

    result = mcp_tools.export_close_package(
        client_id, "2026-01-01", "2026-03-31", str(tmp_path / "legacy-config"))
    assert result["pdf"].startswith(str(tmp_path))
